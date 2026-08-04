"""
Rebuilds luget.db from two Excel source files + four CEFR-level files.

Sources
-------
DeAz   →  ~/Desktop/german/final dics/Dictionary_Last.xlsx
AzDe   →  ~/Desktop/german/final dics/az_de_full_dictionary.xlsx
Levels →  ~/Desktop/german/final dics/levels/Dictionary_{A1,A2,B1,B2}.xlsx
           Level | Id | Word | ... (same columns as DeAz + Level first)

Tables produced
---------------
DeAz             German → Azerbaijani  (+ id + level columns)
AzDe             Azerbaijani → German  (+ id column, read directly from Excel)
bookmark         Empty; schema matches app expectations
quiz_results     Empty; recreated by onOpen as well
training_progress  Empty; ready for the Training screen

user_version = 5  →  bump dbVersion in word_local_data_source_impl.dart to 5.
"""

import shutil
import sqlite3
from pathlib import Path

import openpyxl

BASE        = Path('/Users/elvinmammadov/StudioProjects/Dictionary')
DESKTOP     = Path('/Users/elvinmammadov/Desktop/german/final dics')
EXCEL_DEAZ  = DESKTOP / 'Dictionary_Last.xlsx'
EXCEL_AZDE  = DESKTOP / 'az_de_full_dictionary.xlsx'
LEVELS_DIR  = DESKTOP / 'levels'
DB_OUT      = BASE / 'assets' / 'luget.db'
DB_BACKUP   = DB_OUT.with_suffix('.db.bak')
DB_VERSION  = 5   # bumped: level column + training_progress table added


def _val(row: tuple, col: dict, name: str) -> str | None:
    """Return stripped cell value or None for missing / empty cells."""
    idx = col.get(name)
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


# ---------------------------------------------------------------------------
# 0.  Build id → level mapping from all four level Excel files
# ---------------------------------------------------------------------------

print("Loading CEFR level mappings ...")
id_to_level: dict[str, str] = {}

for level_name in ('A1', 'A2', 'B1', 'B2'):
    path = LEVELS_DIR / f'Dictionary_{level_name}.xlsx'
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    col = {h: i for i, h in enumerate(headers)}
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        word_id = _val(row, col, 'Id')
        if word_id:
            id_to_level[word_id] = level_name
            count += 1
    wb.close()
    print(f"  {level_name}: {count} words")

print(f"  Total mapped IDs: {len(id_to_level)}")


# ---------------------------------------------------------------------------
# 1.  Load DeAz from Dictionary_Last.xlsx  (+ attach level)
# ---------------------------------------------------------------------------

print(f"\nLoading DeAz from {EXCEL_DEAZ} ...")
wb1 = openpyxl.load_workbook(EXCEL_DEAZ, read_only=True, data_only=True)
ws1 = wb1.active
headers1 = list(next(ws1.iter_rows(min_row=1, max_row=1, values_only=True)))
print(f"  Headers: {headers1}")
COL1 = {name: idx for idx, name in enumerate(headers1)}

deaz_rows: list[dict] = []
unlevelled = 0
for row in ws1.iter_rows(min_row=2, values_only=True):
    word        = _val(row, COL1, 'Word')
    translation = _val(row, COL1, 'Translation')
    if not word or not translation:
        continue
    word_id = _val(row, COL1, 'Id')
    level   = id_to_level.get(word_id or '', 'B2')   # default unclassified → B2
    if word_id not in id_to_level:
        unlevelled += 1
    deaz_rows.append({
        'id':          word_id,
        'level':       level,
        'key':         word,
        'value':       translation,
        'article':     _val(row, COL1, 'Article'),
        'gender':      _val(row, COL1, 'Gender'),
        'main_type':   _val(row, COL1, 'MainType'),
        'sub_type':    _val(row, COL1, 'SubType'),
        'genitive':    _val(row, COL1, 'Genitive'),
        'plural':      _val(row, COL1, 'Plural'),
        'imperfekt':   _val(row, COL1, 'Imperfekt'),
        'perfekt':     _val(row, COL1, 'Perfekt'),
        'comparative': _val(row, COL1, 'Comparative'),
        'superlative': _val(row, COL1, 'Superlative'),
        'example':     None,
        'sentence':    _val(row, COL1, 'Sentence'),
    })

wb1.close()
print(f"  DeAz rows loaded   : {len(deaz_rows)}")
print(f"  Rows defaulted B2  : {unlevelled}")

# Level distribution summary
from collections import Counter
dist = Counter(r['level'] for r in deaz_rows)
for lvl in ('A1', 'A2', 'B1', 'B2'):
    print(f"    {lvl}: {dist[lvl]}")


# ---------------------------------------------------------------------------
# 2.  Load AzDe directly from az_de_full_dictionary.xlsx
# ---------------------------------------------------------------------------

print(f"\nLoading AzDe from {EXCEL_AZDE} ...")
wb2 = openpyxl.load_workbook(EXCEL_AZDE, read_only=True, data_only=True)
ws2 = wb2.active
headers2 = list(next(ws2.iter_rows(min_row=1, max_row=1, values_only=True)))
print(f"  Headers: {headers2}")
COL2 = {name: idx for idx, name in enumerate(headers2)}

azde_rows: list[dict] = []
for row in ws2.iter_rows(min_row=2, values_only=True):
    word        = _val(row, COL2, 'Word')
    translation = _val(row, COL2, 'Translation')
    if not word or not translation:
        continue
    azde_rows.append({
        'id':        _val(row, COL2, 'Id'),
        'key':       word,
        'value':     translation,
        'main_type': _val(row, COL2, 'MainType'),
        'sub_type':  _val(row, COL2, 'SubType'),
    })

wb2.close()
print(f"  AzDe rows loaded: {len(azde_rows)}")


# ---------------------------------------------------------------------------
# 3.  Write SQLite database
# ---------------------------------------------------------------------------

if DB_OUT.exists():
    shutil.copy2(DB_OUT, DB_BACKUP)
    print(f"\nBacked up existing DB -> {DB_BACKUP}")
    DB_OUT.unlink()

conn = sqlite3.connect(DB_OUT)
cur  = conn.cursor()

cur.execute(f"PRAGMA user_version = {DB_VERSION}")

cur.executescript("""
CREATE TABLE DeAz (
    id          TEXT,
    level       TEXT,
    key         TEXT,
    value       TEXT,
    article     TEXT,
    gender      TEXT,
    main_type   TEXT,
    sub_type    TEXT,
    genitive    TEXT,
    plural      TEXT,
    imperfekt   TEXT,
    perfekt     TEXT,
    comparative TEXT,
    superlative TEXT,
    example     TEXT,
    sentence    TEXT
);

CREATE INDEX idx_deaz_level ON DeAz(level);
CREATE INDEX idx_deaz_key   ON DeAz(key);

CREATE TABLE AzDe (
    id        TEXT,
    key       TEXT,
    value     TEXT,
    main_type TEXT,
    sub_type  TEXT
);

CREATE INDEX idx_azde_key ON AzDe(key);

CREATE TABLE bookmark (
    key   VARCHAR(100),
    value TEXT,
    date  DATETIME DEFAULT CURRENT_TIMESTAMP,
    type  TEXT,
    PRIMARY KEY (key, value)
);

CREATE TABLE quiz_results (
    id             INTEGER PRIMARY KEY AUTOINCREMENT,
    score          INTEGER NOT NULL,
    totalQuestions INTEGER NOT NULL,
    dateTime       TEXT    NOT NULL
);

CREATE TABLE training_progress (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    word_id     TEXT    NOT NULL,
    level       TEXT    NOT NULL,
    correct     INTEGER NOT NULL DEFAULT 0,
    incorrect   INTEGER NOT NULL DEFAULT 0,
    last_seen   TEXT,
    UNIQUE(word_id)
);
""")

cur.executemany(
    """INSERT INTO DeAz
         (id, level, key, value, article, gender, main_type, sub_type,
          genitive, plural, imperfekt, perfekt, comparative, superlative,
          example, sentence)
       VALUES
         (:id, :level, :key, :value, :article, :gender, :main_type, :sub_type,
          :genitive, :plural, :imperfekt, :perfekt, :comparative, :superlative,
          :example, :sentence)""",
    deaz_rows,
)
print(f"\nInserted {cur.rowcount} rows into DeAz")

cur.executemany(
    """INSERT INTO AzDe (id, key, value, main_type, sub_type)
       VALUES (:id, :key, :value, :main_type, :sub_type)""",
    azde_rows,
)
print(f"Inserted {cur.rowcount} rows into AzDe")

conn.commit()

# Verify
print("\n--- Verification ---")
cur.execute("SELECT COUNT(*) FROM DeAz")
print(f"DeAz total        : {cur.fetchone()[0]}")
for lvl in ('A1', 'A2', 'B1', 'B2'):
    cur.execute("SELECT COUNT(*) FROM DeAz WHERE level = ?", (lvl,))
    print(f"  level={lvl}        : {cur.fetchone()[0]}")
cur.execute("SELECT COUNT(*) FROM AzDe")
print(f"AzDe total        : {cur.fetchone()[0]}")
cur.execute("PRAGMA user_version")
print(f"user_version      : {cur.fetchone()[0]}")

cur.execute("SELECT id, level, key, value FROM DeAz LIMIT 3")
print("\nDeAz sample:")
for r in cur.fetchall():
    print(f"  {r}")

conn.close()
print(f"\nDatabase written to {DB_OUT}")
print("  Bump dbVersion to 5 in word_local_data_source_impl.dart")
